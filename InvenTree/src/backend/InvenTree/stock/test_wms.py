"""Business acceptance for the September 2026 WMS inventory scope."""

import csv
import io
import threading
import uuid
from datetime import datetime, timedelta
from decimal import Decimal
from unittest.mock import patch
from zoneinfo import ZoneInfo

from django.contrib.auth.models import Group, User
from django.contrib.contenttypes.models import ContentType
from django.db import connection, connections
from django.test import TestCase, TransactionTestCase
from django.utils import timezone

from openpyxl import Workbook
from rest_framework.test import APIClient

from common.models import ProjectCode
from common.settings import set_global_setting
from part.models import Part, PartCategory
from users.models import Owner

from .models import (
    ScanStockApprovalRequest,
    ScanStockEvent,
    StockItem,
    StockLocation,
    WmsImportTask,
)
from .services import ApprovalRequestService, ScanStockError, ScanStockService


class WmsAcceptanceTest(TestCase):
    """Exercise public endpoints and native stock side effects together."""

    def setUp(self):
        """Create two projects, two actual locations and an isolated operator."""
        ContentType.objects.clear_cache()
        self.user = User.objects.create_superuser(
            username='wms-user', email='wms@example.test', password='test'
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.category = PartCategory.objects.create(name='设备')
        self.parent = StockLocation.objects.create(name='仓库A')
        self.child = StockLocation.objects.create(name='货架B', parent=self.parent)
        self.part = Part.objects.create(
            name='设备一', IPN='WMS-001', units='count', category=self.category
        )
        self.item = StockItem.objects.create(
            part=self.part,
            location=self.parent,
            batch='P1',
            quantity=10,
            delete_on_deplete=False,
        )
        self.other = StockItem.objects.create(
            part=self.part,
            location=self.child,
            batch='P2',
            quantity=3,
            delete_on_deplete=False,
        )
        set_global_setting('SCAN_APPROVAL_ENABLED', False)

    def test_native_inventory_scope_and_export(self):
        """Native pagination, totals, summary and CSV share every native filter."""
        params = {
            'wms': 'true',
            'location': self.parent.pk,
            'limit': 1,
            'part_detail': 'true',
        }
        result = self.client.get('/api/stock/', params)
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.data['count'], 2)
        self.assertEqual(len(result.data['results']), 1)
        self.assertEqual(result.data['statistics']['units']['count'], Decimal(13))
        self.assertEqual(result.data['results'][0]['wms_equipment']['IPN'], 'WMS-001')
        direct = self.client.get('/api/stock/', {**params, 'cascade': 'false'})
        self.assertEqual(direct.data['count'], 1)
        self.assertEqual(direct.data['statistics']['units']['count'], Decimal(10))
        project = {**params, 'search': 'P2'}
        detail = self.client.get('/api/stock/', project)
        self.assertEqual(detail.data['count'], 1)
        self.assertEqual(detail.data['statistics']['units']['count'], Decimal(3))
        summary = self.client.get('/api/stock/', {**project, 'view': 'summary'})
        self.assertEqual(Decimal(summary.data['results'][0]['quantity']), Decimal(3))
        self.assertEqual(summary.data['results'][0]['project_count'], 1)
        export = self.client.get('/api/stock/', {**project, 'export': 'csv'})
        self.assertEqual(export.status_code, 200)
        self.assertIn('P2', export.content.decode('utf-8-sig'))
        self.assertNotIn('P1', export.content.decode('utf-8-sig'))
        empty = self.client.get('/api/stock/', {**params, 'batch': 'missing'})
        self.assertEqual(empty.data['count'], 0)
        self.assertEqual(empty.data['statistics']['part_count'], 0)

    def test_native_inventory_zero_and_catalogue_filters(self):
        """Native WMS filters retain zero items only when explicitly requested."""
        self.part.brand = 'NativeBrand'
        self.part.save()
        self.item.quantity = 0
        self.item.save()
        params = {'wms': 'true', 'brand': 'NativeBrand'}
        default = self.client.get('/api/stock/', params)
        self.assertEqual(default.data['count'], 1)
        zeros = self.client.get('/api/stock/', {**params, 'include_zero': 'true'})
        self.assertEqual(zeros.data['count'], 2)
        empty = self.client.get('/api/stock/', {**params, 'brand': 'Other'})
        self.assertEqual(empty.data['count'], 0)
        search = self.client.get(
            '/api/stock/', {'wms': 'true', 'search': 'NativeBrand'}
        )
        self.assertEqual(search.data['count'], 1)
        frozen = self.client.get('/api/stock/', {**params, 'has_frozen': 'true'})
        self.assertEqual(frozen.data['count'], 0)

    def test_native_export_import_roundtrip(self):
        """The downloaded CSV can be uploaded unchanged without changing balances."""
        exported = self.client.get(
            '/api/stock/', {'wms': 'true', 'export': 'csv', 'limit': 1}
        )
        template = self.client.get('/api/stock/wms/imports/', {'kind': 'count'})
        self.assertEqual(
            next(csv.reader(io.StringIO(exported.content.decode('utf-8-sig')))),
            next(csv.reader(io.StringIO(template.content.decode('utf-8-sig')))),
        )
        self.assertIn(
            '项目', next(csv.reader(io.StringIO(template.content.decode('utf-8-sig'))))
        )
        upload = io.BytesIO(exported.content)
        upload.name = 'inventory.csv'
        preview = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'count', 'file': upload},
            format='multipart',
        )
        self.assertEqual(preview.status_code, 201, preview.data)
        self.assertTrue(preview.data['can_commit'])
        self.assertEqual(len(preview.data['results']), 2)
        self.assertTrue(
            all(Decimal(row['delta']) == 0 for row in preview.data['results'])
        )
        self.post('wms/imports', {'task': preview.data['task']})
        self.item.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)
        self.assertEqual(self.other.quantity, 3)

    def test_edited_export_xlsx_import(self):
        """Editing quantity in an exported sheet changes only that scoped balance."""
        exported = self.client.get(
            '/api/stock/', {'wms': 'true', 'export': 'csv', 'batch': 'P1'}
        )
        rows = list(csv.reader(io.StringIO(exported.content.decode('utf-8-sig'))))
        rows[1][rows[0].index('盘点数量')] = '8'
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        upload = io.BytesIO()
        workbook.save(upload)
        upload.seek(0)
        upload.name = 'inventory.xlsx'
        preview = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'count', 'file': upload},
            format='multipart',
        )
        self.assertEqual(preview.status_code, 201, preview.data)
        self.assertTrue(preview.data['can_commit'])
        self.assertEqual(Decimal(preview.data['results'][0]['delta']), -2)
        self.post('wms/imports', {'task': preview.data['task']})
        self.item.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual(self.item.quantity, 8)
        self.assertEqual(self.other.quantity, 3)

    def test_legacy_count_quantity_header(self):
        """Previously supplied count templates remain valid."""
        upload = io.BytesIO(
            f'设备ID,项目,地点完整路径,盘点数量\nWMS-001,P1,{self.parent.pathstring},8\n'.encode()
        )
        upload.name = 'legacy.csv'
        response = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'count', 'file': upload},
            format='multipart',
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(response.data['can_commit'])
        self.assertEqual(Decimal(response.data['results'][0]['delta']), -2)

    def test_count_import_reuses_unique_project_archive(self):
        """Preview is read-only; imported stock shares one project and retries do not recount."""
        project = '00123'
        preview = self.preview(
            [
                self.count_row(batch=project, quantity='2'),
                self.count_row(
                    batch=project, location=self.child.pathstring, quantity='3'
                ),
            ]
        )
        self.assertTrue(preview['can_commit'], preview)
        self.assertFalse(ProjectCode.objects.filter(code=project).exists())
        result = self.post('wms/imports', {'task': preview['task']})
        archive = ProjectCode.objects.get(code=project)
        items = StockItem.objects.filter(batch=project).order_by('quantity')
        self.assertEqual(list(items.values_list('quantity', flat=True)), [2, 3])
        self.assertEqual(
            set(items.values_list('project_code_id', flat=True)), {archive.pk}
        )
        replay = self.post('wms/imports', {'task': preview['task']})
        self.assertTrue(replay['replayed'])
        self.assertEqual(replay['results'], result['results'])
        self.assertEqual(ProjectCode.objects.filter(code=project).count(), 1)
        self.assertFalse(ScanStockEvent.objects.exists())

    def test_count_project_length_rejected_before_commit(self):
        """An oversized project must be a row error instead of a database failure."""
        preview = self.preview([self.count_row(batch='项' * 101)])
        self.assertFalse(preview['can_commit'], preview)
        self.assertIn('error', preview['results'][0])
        self.post('wms/imports', {'task': preview['task']}, 400)
        self.assertEqual(StockItem.objects.count(), 2)
        self.assertEqual(ProjectCode.objects.count(), 2)

    def test_equipment_csv_create_and_xlsx_update_without_stock(self):
        """Both file formats preserve textual IDs and import catalogue data only."""
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(['设备ID', '设备名称', '品牌', '设备分类完整路径', '单位'])
        writer.writerow(
            ['00123', '中文设备', '原品牌', self.category.pathstring, 'count']
        )
        writer.writerow(['', '自动编号设备', '', self.category.pathstring, 'count'])
        upload = io.BytesIO(stream.getvalue().encode('utf-8-sig'))
        upload.name = 'parts.csv'
        preview = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'parts', 'file': upload},
            format='multipart',
        )
        self.assertEqual(preview.status_code, 201, preview.data)
        self.assertTrue(preview.data['can_commit'], preview.data)
        self.assertEqual(Part.objects.count(), 1)
        self.post('wms/imports', {'task': preview.data['task']})
        self.assertEqual(Part.objects.count(), 3)
        automatic = Part.objects.get(name='自动编号设备')
        self.assertTrue(automatic.IPN)
        part = Part.objects.get(IPN='00123')
        workbook = Workbook()
        workbook.active.append(['设备ID', '设备名称', '品牌', '型号', '规格', '备注'])
        workbook.active.append(
            ['00123', '更新名称', '', '型号甲', '规格甲', '导入备注']
        )
        upload = io.BytesIO()
        workbook.save(upload)
        upload.seek(0)
        upload.name = 'parts.xlsx'
        preview = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'parts', 'file': upload},
            format='multipart',
        )
        self.assertEqual(preview.status_code, 201, preview.data)
        self.assertTrue(preview.data['can_commit'], preview.data)
        self.post('wms/imports', {'task': preview.data['task']})
        part.refresh_from_db()
        self.assertEqual((part.name, part.brand), ('更新名称', '原品牌'))
        self.assertEqual(
            (part.model_number, part.specification, part.notes),
            ('型号甲', '规格甲', '导入备注'),
        )
        self.assertEqual(Part.objects.count(), 3)
        self.assertEqual(
            list(StockItem.objects.order_by('pk').values_list('quantity', flat=True)),
            [10, 3],
        )

    def test_equipment_duplicate_file_rows_block_entire_import(self):
        """A valid new row cannot be saved when another row repeats a device identity."""
        upload = io.BytesIO(
            (
                '设备ID,设备名称,设备分类完整路径,单位\n'
                f'NEW-IMPORT,新设备,{self.category.pathstring},count\n'
                f'WMS-001,更新设备,{self.category.pathstring},count\n'
                f'wms-001,重复设备,{self.category.pathstring},count\n'
            ).encode('utf-8-sig')
        )
        upload.name = 'duplicate-parts.csv'
        preview = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'parts', 'file': upload},
            format='multipart',
        )
        self.assertEqual(preview.status_code, 201, preview.data)
        self.assertFalse(preview.data['can_commit'])
        self.assertIn('重复设备', preview.data['results'][2]['error'])
        self.post('wms/imports', {'task': preview.data['task']}, 400)
        self.assertFalse(Part.objects.filter(IPN='NEW-IMPORT').exists())
        self.part.refresh_from_db()
        self.assertEqual(self.part.name, '设备一')

    def post(self, resource, data, code=200):
        """Post JSON and report the business error on unexpected responses."""
        result = self.client.post(f'/api/stock/{resource}/', data, format='json')
        self.assertEqual(result.status_code, code, result.data)
        return result.data

    def count_row(self, **extra):
        """Build one valid direct-location count row."""
        return {
            'IPN': self.part.IPN,
            'batch': 'P1',
            'location': self.parent.pathstring,
            'quantity': '8',
            **extra,
        }

    def preview(self, rows=None, kind='count'):
        """Create a persisted preview task."""
        return self.post(
            'wms/imports', {'kind': kind, 'rows': rows or [self.count_row()]}, 201
        )

    def scan(self, operation='inbound', **extra):
        """Build and execute one independent scan."""
        data = {
            'barcode': self.part.IPN,
            'client_event_id': str(uuid.uuid4()),
            'quantity': '1',
            'batch': 'P1',
            'purpose': '安装',
            (
                'destination_location' if operation == 'inbound' else 'source_location'
            ): self.parent.pk,
            **extra,
        }
        return self.client.post(f'/api/stock/scan/{operation}/', data, format='json')

    def approval(self, direction='OUTBOUND'):
        """Configure an existing single-stage route and create its draft."""
        set_global_setting('SCAN_APPROVAL_ENABLED', True)
        approver = User.objects.create_superuser(
            username='wms-approver', email='a@example.test', password='test'
        )
        group = Group.objects.create(name='WMS审批')
        approver.groups.add(group)
        self.parent.approval_owner = Owner.get_owner(group)
        self.parent.save()
        request = ApprovalRequestService.create_request(
            user=self.user, direction=direction, location=self.parent
        )
        return request, approver

    def options(self, **params):
        """Read outbound scope without mutating stock."""
        response = self.client.get(
            '/api/stock/wms/options/',
            {
                'outbound': 'true',
                'barcode': self.part.IPN,
                'location': self.parent.pk,
                **params,
            },
        )
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def test_options_exact_scope_and_zero(self):
        """Retained zero, missing rows and child stock do not hide scope state."""
        self.item.quantity = 0
        self.item.save()
        data = self.options()
        self.assertEqual(data['projects'], [])
        self.assertEqual(data['stock_status']['state'], 'NO_STOCK')
        self.assertEqual(Decimal(data['stock_status']['quantity']), 0)
        self.assertEqual(data['stock_status']['scope'], 'all_projects')
        self.assertEqual(self.options(batch='P2')['stock_status']['state'], 'NO_STOCK')
        self.item.delete()
        self.assertEqual(self.options()['stock_status']['state'], 'NO_STOCK')

    def test_options_scan_preview_without_stock_effects(self):
        """Equipment and stock codes expose device identity before a confirmed scan."""
        for code in (self.part.IPN, f'INV-PA{self.part.pk}', f'INV-SI{self.item.pk}'):
            with self.subTest(code=code):
                data = self.options(barcode=code)
                self.assertEqual(
                    data['scanned_part'],
                    {'pk': self.part.pk, 'IPN': self.part.IPN, 'name': self.part.name},
                )
        inbound = self.client.get('/api/stock/wms/options/', {'barcode': self.part.IPN})
        self.assertEqual(inbound.status_code, 200)
        self.assertEqual(inbound.data['scanned_part']['pk'], self.part.pk)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)
        self.assertEqual(ScanStockEvent.objects.count(), 0)
        self.assertEqual(ScanStockApprovalRequest.objects.count(), 0)
        self.assertIsNone(self.options(barcode='unknown-device')['scanned_part'])

    def test_options_invalid_and_incomplete_are_not_zero(self):
        """Bad codes, wrong object kinds and incomplete conditions stay distinct."""
        for barcode in ('unknown', 'INV-PA-bad', f'INV-SL{self.parent.pk}'):
            with self.subTest(barcode=barcode):
                status = self.options(barcode=barcode)['stock_status']
                self.assertEqual(status['state'], 'INVALID_BARCODE')
                self.assertNotIn('quantity', status)
        self.assertEqual(
            self.options(barcode='')['stock_status']['state'], 'INCOMPLETE'
        )
        self.assertEqual(
            self.options(location='bad')['stock_status']['state'], 'INVALID_LOCATION'
        )

    def test_options_stock_code_checks_scope_and_aggregates_siblings(self):
        """Item codes validate context but balances cover the entire combination."""
        StockItem.objects.create(
            part=self.part, location=self.parent, batch='P1', quantity=2
        )
        code = f'INV-SI{self.item.pk}'
        status = self.options(barcode=code, batch='P1')['stock_status']
        self.assertEqual(Decimal(status['quantity']), 12)
        for params in ({'batch': 'P2'}, {'location': self.child.pk}):
            self.assertEqual(
                self.options(barcode=code, **params)['stock_status']['state'],
                'CONTEXT_MISMATCH',
            )
        self.assertEqual(
            self.options(barcode=f'INV-PA{self.part.pk}')['projects'], ['P1']
        )

    def test_options_freezes_and_approval_posting(self):
        """Reservations change availability; only approval changes physical stock."""
        request, approver = self.approval()
        self.assertEqual(
            self.scan(
                'outbound', quantity=10, approval_request_id=request.pk
            ).status_code,
            201,
        )
        status = self.options(batch='P1')['stock_status']
        self.assertEqual(status['state'], 'NO_AVAILABLE_STOCK')
        self.assertEqual(Decimal(status['quantity']), 10)
        self.assertEqual(Decimal(status['frozen']), 10)
        ApprovalRequestService.submit(request_id=request.pk, user=self.user)
        ApprovalRequestService.approve(request_id=request.pk, user=approver)
        self.assertEqual(self.options(batch='P1')['stock_status']['state'], 'NO_STOCK')

    def test_options_native_allocation_and_release(self):
        """Native occupancy and scan freezes use existing availability rules."""
        from company.models import Company
        from order.models import SalesOrder, SalesOrderAllocation, SalesOrderLineItem

        customer = Company.objects.create(name='status-customer', is_customer=True)
        order = SalesOrder.objects.create(reference='WMS-STATUS-SO', customer=customer)
        line = SalesOrderLineItem.objects.create(
            order=order, part=self.part, quantity=10
        )
        allocation = SalesOrderAllocation.objects.create(
            line=line, item=self.item, quantity=10
        )
        status = self.options()['stock_status']
        self.assertEqual(status['state'], 'NO_AVAILABLE_STOCK')
        self.assertEqual(Decimal(status['allocated']), 10)
        allocation.delete()
        self.assertEqual(self.options()['stock_status']['state'], 'AVAILABLE')
        SalesOrderAllocation.objects.create(line=line, item=self.item, quantity=4)
        self.assertEqual(Decimal(self.options()['stock_status']['available']), 6)
        self.assertEqual(self.scan('outbound', quantity=7).status_code, 400)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)

    def test_options_decimal_depletion_and_multiple_rows(self):
        """One depleted item does not mean the project is empty; decimals stay exact."""
        self.item.quantity = Decimal('0.00001')
        self.item.save()
        sibling = StockItem.objects.create(
            part=self.part,
            location=self.parent,
            batch='P1',
            quantity=Decimal('0.00002'),
        )
        self.assertEqual(
            self.scan(
                'outbound', quantity='0.00001', stock_item_id=self.item.pk
            ).status_code,
            201,
        )
        status = self.options(batch='P1')['stock_status']
        self.assertEqual(status['state'], 'AVAILABLE')
        self.assertEqual(Decimal(status['quantity']), Decimal('0.00002'))
        self.assertEqual(
            self.scan(
                'outbound', quantity='0.00002', stock_item_id=sibling.pk
            ).status_code,
            201,
        )
        self.assertEqual(self.options(batch='P1')['stock_status']['state'], 'NO_STOCK')

    def test_options_new_projects_only_after_posting(self):
        """New names are deduplicated and appear only once stock is saved."""
        request, approver = self.approval('INBOUND')
        self.assertEqual(
            self.scan(batch='NEW', approval_request_id=request.pk).status_code, 201
        )
        self.assertNotIn('NEW', self.options(outbound='false')['projects'])
        ApprovalRequestService.submit(request_id=request.pk, user=self.user)
        self.assertNotIn('NEW', self.options(outbound='false')['projects'])
        ApprovalRequestService.approve(request_id=request.pk, user=approver)
        self.assertIn('NEW', self.options(outbound='false')['projects'])
        self.assertEqual(self.options()['projects'].count('P1'), 1)

    def test_options_do_not_mix_equipment_or_empty_projects(self):
        """Candidates are saved, distinct, nonempty names in the exact scope."""
        other_part = Part.objects.create(
            name='Other device', IPN='OTHER-DEVICE', units='count'
        )
        StockItem.objects.create(
            part=other_part, location=self.parent, batch='OTHER', quantity=5
        )
        StockItem.objects.create(
            part=self.part, location=self.parent, batch='', quantity=5
        )
        StockItem.objects.create(
            part=self.part, location=self.parent, batch='P1', quantity=2
        )
        self.assertEqual(self.options()['projects'], ['P1'])
        imported = self.preview([self.count_row(batch='IMPORTED', quantity='2')])
        self.assertNotIn('IMPORTED', self.options()['projects'])
        self.post('wms/imports', {'task': imported['task']})
        self.assertIn('IMPORTED', self.options()['projects'])

    def test_options_after_outbound_can_have_stock_but_no_availability(self):
        """Posting can leave physical stock fully occupied by a native order."""
        from company.models import Company
        from order.models import SalesOrder, SalesOrderAllocation, SalesOrderLineItem

        sibling = StockItem.objects.create(
            part=self.part, location=self.parent, batch='P1', quantity=2
        )
        customer = Company.objects.create(name='remaining-customer', is_customer=True)
        order = SalesOrder.objects.create(
            reference='WMS-REMAINING-SO', customer=customer
        )
        line = SalesOrderLineItem.objects.create(
            order=order, part=self.part, quantity=2
        )
        SalesOrderAllocation.objects.create(line=line, item=sibling, quantity=2)
        self.assertEqual(
            self.scan('outbound', quantity=10, stock_item_id=self.item.pk).status_code,
            201,
        )
        status = self.options(batch='P1')['stock_status']
        self.assertEqual(status['state'], 'NO_AVAILABLE_STOCK')
        self.assertEqual(Decimal(status['quantity']), 2)
        self.assertEqual(Decimal(status['available']), 0)

    def test_project_and_location_summary(self):
        """Parent totals contain each native row once and filters stay scoped."""

        def quantity(**params):
            response = self.client.get(
                '/api/stock/wms/inventory/',
                {'view': 'summary', 'location': self.parent.pk, **params},
            )
            self.assertEqual(response.status_code, 200, response.data)
            return response.data['results'][0]

        self.assertEqual(Decimal(quantity()['quantity']), 13)
        self.assertEqual(quantity()['project_count'], 2)
        self.assertEqual(Decimal(quantity(include_children='false')['quantity']), 10)
        self.assertEqual(Decimal(quantity(batch='P2')['quantity']), 3)
        self.assertEqual(quantity(batch='P2')['project_count'], 1)

    def test_scan_project_and_purpose_required(self):
        """Blank project, missing purpose and conflicting aliases fail before stock."""
        for response in [
            self.scan(batch=' '),
            self.scan('outbound', purpose=''),
            self.scan(project_number='P2'),
        ]:
            self.assertEqual(response.status_code, 400, response.data)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)

    def test_alias_and_compatible_inbound_reuse(self):
        """Trimmed aliases use the same project balance without creating duplicates."""
        result = self.scan(batch=' P1 ', project_number='P1')
        self.assertEqual(result.status_code, 201, result.data)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 11)
        self.assertEqual(
            StockItem.objects.filter(part=self.part, location=self.parent).count(), 1
        )

    def test_outbound_does_not_borrow(self):
        """A different project or child location cannot satisfy an outbound request."""
        for data in [
            {'batch': 'P2'},
            {'quantity': 11},
            {'stock_item_id': self.other.pk},
        ]:
            result = self.scan('outbound', **data)
            self.assertFalse(result.data.get('success', False), result.data)
        self.item.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual((self.item.quantity, self.other.quantity), (10, 3))

    def test_idempotency_includes_purpose_and_remark(self):
        """Transport retry replays once while separate scans accumulate normally."""
        event_id = str(uuid.uuid4())
        first = self.scan(client_event_id=event_id, remark='备注')
        self.assertEqual(first.status_code, 201, first.data)
        retry = self.scan(client_event_id=event_id, remark='备注')
        self.assertTrue(retry.data['replayed'])
        for field in [{'remark': '不同'}, {'purpose': '其他'}, {'batch': 'P2'}]:
            conflict = self.scan(client_event_id=event_id, **field)
            self.assertEqual(conflict.status_code, 409, conflict.data)
        self.scan()
        self.scan()
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 13)

    def test_count_absolute_not_scan_outbound_and_retry(self):
        """10 to 8 is a -2 stocktake, leaves other stock alone and safely replays."""
        preview = self.preview()
        self.assertTrue(preview['can_commit'], preview)
        self.assertEqual(Decimal(preview['results'][0]['delta']), -2)
        first = self.post('wms/imports', {'task': preview['task']})
        second = self.post('wms/imports', {'task': preview['task']})
        self.assertTrue(second['replayed'])
        self.assertEqual(first['results'], second['results'])
        self.item.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual((self.item.quantity, self.other.quantity), (8, 3))
        self.assertEqual(ScanStockEvent.objects.count(), 0)
        self.assertIn('盘点任务', self.item.tracking_info.latest('pk').notes)

    def test_invalid_count_rows_block_whole_batch(self):
        """Unknown paths, duplicates and illegal precision block all mutations."""
        for extra in [
            {'batch': ''},
            {'location': '不存在'},
            {'quantity': '-1'},
            {'quantity': 'NaN'},
            {'quantity': '0.000001'},
            {'quantity': '10000000000'},
            {'units': 'm'},
            {'name': '错误名称'},
            {'count_date': 'not-a-date'},
        ]:
            preview = self.preview([self.count_row(**extra)])
            self.assertFalse(preview['can_commit'], preview)
            self.post('wms/imports', {'task': preview['task']}, 400)
        preview = self.preview([self.count_row(), self.count_row()])
        self.assertFalse(preview['can_commit'])
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)

    def test_ambiguous_stock_rejected(self):
        """Counts never distribute or overwrite ambiguous native stock items."""
        StockItem.objects.create(
            part=self.part, location=self.parent, batch='P1', quantity=2
        )
        preview = self.preview()
        self.assertFalse(preview['can_commit'])
        self.assertIn('多条', preview['results'][0]['error'])

    def test_stale_quantity_and_structure(self):
        """Changed target quantity or project prevents committing an old preview."""
        preview = self.preview()
        self.item.add_stock(1, self.user)
        self.post('wms/imports', {'task': preview['task']}, 400)
        preview = self.preview()
        self.item.batch = 'P3'
        self.item.save()
        self.post('wms/imports', {'task': preview['task']}, 400)

    def test_freeze_blocks_count_and_stale_preview(self):
        """Physical stock stays unchanged in draft and counts cannot erase freezes."""
        preview = self.preview()
        request, _ = self.approval()
        result = self.scan('outbound', quantity=8, approval_request_id=request.pk)
        self.assertEqual(result.status_code, 201, result.data)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)
        self.assertEqual(self.item.scan_frozen_quantity(), 8)
        failed = self.preview([self.count_row(quantity='5')])
        self.assertFalse(failed['can_commit'])
        self.assertIn(request.reference, failed['results'][0]['error'])
        self.post('wms/imports', {'task': preview['task']}, 400)
        ApprovalRequestService.cancel(request_id=request.pk, user=self.user)
        self.assertEqual(self.item.scan_frozen_quantity(), 0)

    def test_approval_inbound_defers_stock_and_snapshots(self):
        """Inbound approval records requester separately from posting approver."""
        request, approver = self.approval('INBOUND')
        result = self.scan(approval_request_id=request.pk, remark='入库备注')
        self.assertEqual(result.status_code, 201, result.data)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)
        ApprovalRequestService.submit(request_id=request.pk, user=self.user)
        ApprovalRequestService.approve(request_id=request.pk, user=approver)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 11)
        event = ScanStockEvent.objects.get(pk=result.data['event']['pk'])
        self.assertEqual(event.operator_username_snapshot, self.user.username)
        self.assertEqual(event.approver_snapshot, approver.username)
        self.assertEqual(
            (event.units_snapshot, event.batch_snapshot, event.remark),
            ('count', 'P1', '入库备注'),
        )

    def test_approval_unknown_equipment_deferred(self):
        """Unknown equipment and stock are created only after whole-request approval."""
        request, approver = self.approval('INBOUND')
        result = self.scan(
            barcode='NEW-PART',
            approval_request_id=request.pk,
            create_part={
                'name': '新设备',
                'units': 'm',
                'category': self.category.pk,
                'default_location': self.parent.pk,
            },
        )
        self.assertEqual(result.status_code, 201, result.data)
        self.assertFalse(Part.objects.filter(IPN='NEW-PART').exists())
        ApprovalRequestService.submit(request_id=request.pk, user=self.user)
        ApprovalRequestService.approve(request_id=request.pk, user=approver)
        part = Part.objects.get(IPN='NEW-PART')
        self.assertEqual(part.units, 'm')
        self.assertEqual(part.stock_items.get().batch, 'P1')

    def test_whole_approval_failure_rolls_back(self):
        """A later apply failure retains all freezes with no partial postings."""
        request, approver = self.approval('OUTBOUND')
        for _ in range(2):
            self.assertEqual(
                self.scan('outbound', approval_request_id=request.pk).status_code, 201
            )
        ApprovalRequestService.submit(request_id=request.pk, user=self.user)
        original = ScanStockService._apply_outbound
        calls = []

        def fail_second(service, event, resolved):
            calls.append(event.pk)
            if len(calls) == 2:
                raise ScanStockError('TEST_FAILURE', '验收模拟第二行失败')
            return original(service, event, resolved)

        with patch.object(ScanStockService, '_apply_outbound', fail_second):
            with self.assertRaises(ScanStockError):
                ApprovalRequestService.approve(request_id=request.pk, user=approver)
        request.refresh_from_db()
        self.assertTrue(request.last_approval_error)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)
        self.assertEqual(self.item.scan_frozen_quantity(), 2)
        self.assertFalse(request.lines.filter(outcome='SUCCESS').exists())
        ApprovalRequestService.approve(request_id=request.pk, user=approver)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 8)
        self.assertEqual(self.item.scan_frozen_quantity(), 0)

    def test_count_commit_rolls_back_later_failure(self):
        """A second-row failure rolls back the first adjustment and task result."""
        preview = self.preview(
            [
                self.count_row(),
                self.count_row(
                    batch='P2', location=self.child.pathstring, quantity='1'
                ),
            ]
        )
        original = StockItem.stocktake

        def fail_second(item, *args, **kwargs):
            if item.pk == self.other.pk:
                from rest_framework.exceptions import ValidationError

                raise ValidationError('模拟盘点失败')
            return original(item, *args, **kwargs)

        with patch.object(StockItem, 'stocktake', fail_second):
            self.post('wms/imports', {'task': preview['task']}, 400)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)
        self.assertIsNone(WmsImportTask.objects.get(pk=preview['task']).completed_at)

    def test_count_new_equipment_once_multiple_projects(self):
        """One new master is shared across project rows; conflicting data is rejected."""
        row = self.count_row(
            IPN='NEW', name='新设备', units='m', category=self.category.pathstring
        )
        preview = self.preview([row, {**row, 'batch': 'P2'}])
        self.assertTrue(preview['can_commit'], preview)
        self.post('wms/imports', {'task': preview['task']})
        self.assertEqual(Part.objects.filter(IPN='NEW').count(), 1)
        self.assertEqual(StockItem.objects.filter(part__IPN='NEW').count(), 2)
        conflict = self.preview(
            [
                {**row, 'IPN': 'OTHER'},
                {**row, 'IPN': 'OTHER', 'batch': 'P2', 'units': 'count'},
            ]
        )
        self.assertFalse(conflict['can_commit'])

    def test_equipment_import_preserves_optional_and_detects_stale(self):
        """Blank optional fields preserve values and changed masters invalidate preview."""
        self.part.brand = '原品牌'
        self.part.save()
        preview = self.preview(
            [{'IPN': 'wms-001', 'brand': '', 'name': '新名称'}], kind='parts'
        )
        self.assertTrue(preview['can_commit'], preview)
        self.post('wms/imports', {'task': preview['task']})
        self.part.refresh_from_db()
        self.assertEqual(self.part.brand, '原品牌')
        preview = self.preview([{'IPN': 'WMS-001', 'brand': '新品牌'}], kind='parts')
        self.part.name = '其他人修改'
        self.part.save()
        self.post('wms/imports', {'task': preview['task']}, 400)
        duplicate = self.preview(
            [{'IPN': 'WMS-001'}, {'IPN': ' wms-001 '}], kind='parts'
        )
        self.assertFalse(duplicate['can_commit'])

    def test_equipment_units_and_id_protected(self):
        """WMS editing and import cannot change identities with existing business."""
        self.post('wms/parts', {'pk': self.part.pk, 'IPN': 'CHANGED'}, 400)
        preview = self.preview([{'IPN': self.part.IPN, 'units': 'm'}], kind='parts')
        self.assertFalse(preview['can_commit'])

    def test_posting_date_and_units_all_pages(self):
        """China end-of-day is included; failed/pending scans and next day are excluded."""
        events = []
        for _ in range(4):
            response = self.scan()
            self.assertEqual(response.status_code, 201, response.data)
            events.append(response.data['event']['pk'])
        china = ZoneInfo('Asia/Shanghai')
        ScanStockEvent.objects.filter(pk=events[0]).update(
            completed_at=datetime(2026, 9, 7, 23, 59, 59, tzinfo=china)
        )
        ScanStockEvent.objects.filter(pk=events[1]).update(
            completed_at=datetime(2026, 9, 8, 0, 0, tzinfo=china)
        )
        ScanStockEvent.objects.filter(pk=events[2]).update(
            completed_at=datetime(2026, 9, 1, 0, 0, tzinfo=china), units_snapshot='m'
        )
        ScanStockEvent.objects.filter(pk=events[3]).update(
            outcome='FAILED',
            applied_quantity=0,
            completed_at=datetime(2026, 9, 7, 12, 0, tzinfo=china),
        )
        response = self.client.get(
            '/api/stock/wms/records/',
            {'start': '2026-09-01', 'end': '2026-09-07', 'limit': 1},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['count'], 2)
        self.assertEqual(response.data['statistics']['record_count'], 2)
        self.assertEqual(
            {k: Decimal(v) for k, v in response.data['statistics']['units'].items()},
            {'count': 1, 'm': 1},
        )
        self.part.name = '改名'
        self.part.active = False
        self.part.save()
        result = self.client.get(
            '/api/stock/wms/records/', {'start': '2026-09-07', 'end': '2026-09-07'}
        ).data['results'][0]
        self.assertEqual(result['name'], '设备一')

    def test_zero_hidden_export_scope_and_notes(self):
        """Exports use the filtered full set and preserve separate inventory notes."""
        self.item.notes = '库存备注'
        self.item.save()
        self.post(
            'wms/imports',
            {
                'task': self.preview([self.count_row(quantity='0', remark='盘点备注')])[
                    'task'
                ]
            },
        )
        response = self.client.get('/api/stock/wms/inventory/', {'batch': 'P1'})
        self.assertEqual(response.data['count'], 0)
        response = self.client.get(
            '/api/stock/wms/inventory/',
            {'batch': 'P1', 'include_zero': 'true', 'view': 'detail'},
        )
        self.assertEqual(response.data['count'], 1)
        self.assertIn('库存备注', response.data['results'][0]['stock_notes'][0])
        export = self.client.get(
            '/api/stock/wms/inventory/', {'batch': 'P2', 'export': 'csv', 'limit': 1}
        )
        content = export.content.decode('utf-8-sig')
        self.assertIn('P2', content)
        self.assertNotIn('P1', content)

    def test_xlsx_upload_and_permissions(self):
        """Real XLSX uploads share validation and unprivileged readers cannot export."""
        workbook = Workbook()
        workbook.active.append(['设备ID', '项目', '地点完整路径', '盘点数量'])
        workbook.active.append([self.part.IPN, 'P1', self.parent.pathstring, 8])
        file = io.BytesIO()
        workbook.save(file)
        file.seek(0)
        file.name = 'count.xlsx'
        response = self.client.post(
            '/api/stock/wms/imports/',
            {'kind': 'count', 'file': file},
            format='multipart',
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(response.data['can_commit'])
        other = User.objects.create_user(username='no-role')
        self.client.force_authenticate(other)
        for resource in ('inventory', 'records', 'parts', 'options'):
            response = self.client.get(f'/api/stock/wms/{resource}/', {'export': 'csv'})
            self.assertEqual(response.status_code, 403)
        self.post('wms/imports', {'task': response.data.get('task', 1)}, 400)

    def test_reject_expiry_and_remove_release_freezes(self):
        """All existing terminal paths release reservations without deducting stock."""
        request, approver = self.approval()
        result = self.scan('outbound', approval_request_id=request.pk)
        event_id = result.data['event']['pk']
        ApprovalRequestService.remove_line(
            request_id=request.pk, event_id=event_id, user=self.user
        )
        self.assertEqual(self.item.scan_frozen_quantity(), 0)
        self.scan('outbound', approval_request_id=request.pk)
        ApprovalRequestService.submit(request_id=request.pk, user=self.user)
        with self.assertRaises(ScanStockError):
            ApprovalRequestService.approve(request_id=request.pk, user=self.user)
        ApprovalRequestService.reject(
            request_id=request.pk, user=approver, note='验收驳回'
        )
        self.assertEqual(self.item.scan_frozen_quantity(), 0)
        request = ApprovalRequestService.create_request(
            user=self.user, direction='OUTBOUND', location=self.parent
        )
        self.scan('outbound', approval_request_id=request.pk)
        ScanStockApprovalRequest.objects.filter(pk=request.pk).update(
            expires_at=timezone.now() - timedelta(minutes=1)
        )
        self.assertTrue(ApprovalRequestService.expire(request.pk))
        self.assertEqual(self.item.scan_frozen_quantity(), 0)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)

    def test_required_master_fields_cannot_be_cleared(self):
        """The WMS interface enforces required fields during both create and edit."""
        for field in ('IPN', 'name', 'units', 'category'):
            self.post('wms/parts', {'pk': self.part.pk, field: ''}, 400)
        self.post(
            'wms/parts',
            {
                'IPN': 'NEW-NO-UNIT',
                'name': 'missing unit',
                'category': self.category.pathstring,
            },
            400,
        )
        result = self.scan(
            barcode='NEW-SETUP',
            create_part={
                'name': 'new',
                'category': self.category.pk,
                'default_location': self.parent.pk,
            },
        )
        self.assertEqual(result.status_code, 400)

    def test_zero_can_be_recounted_without_master_edit_changing_timestamp(self):
        """Retained zero balances remain valid count targets; catalogue edits do not move them."""
        self.item.stocktake(0, self.user)
        self.item.refresh_from_db()
        changed_at = self.item.inventory_changed_at
        self.part.brand = 'new-brand'
        self.part.save()
        self.item.refresh_from_db()
        self.assertEqual(changed_at, self.item.inventory_changed_at)
        preview = self.preview()
        self.assertTrue(preview['can_commit'], preview)
        self.post('wms/imports', {'task': preview['task']})
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 8)

    def test_whole_import_does_not_leave_new_master(self):
        """A failure after creating a new master rolls back every creation and count."""
        from rest_framework.exceptions import ValidationError

        rows = [
            self.count_row(
                IPN='ROLLBACK-NEW',
                name='new',
                units='m',
                category=self.category.pathstring,
            ),
            self.count_row(),
        ]
        preview = self.preview(rows)
        original = StockItem.stocktake

        def fail_existing(item, *args, **kwargs):
            if item.pk == self.item.pk:
                raise ValidationError('later failure')
            return original(item, *args, **kwargs)

        with patch.object(StockItem, 'stocktake', fail_existing):
            self.post('wms/imports', {'task': preview['task']}, 400)
        self.assertFalse(Part.objects.filter(IPN='ROLLBACK-NEW').exists())

    def test_native_allocation_prevents_count_below_occupied(self):
        """Native sales allocations remain protected during count imports."""
        from company.models import Company
        from order.models import SalesOrder, SalesOrderAllocation, SalesOrderLineItem

        customer = Company.objects.create(name='count-customer', is_customer=True)
        order = SalesOrder.objects.create(reference='WMS-COUNT-SO', customer=customer)
        line = SalesOrderLineItem.objects.create(
            order=order, part=self.part, quantity=9
        )
        SalesOrderAllocation.objects.create(line=line, item=self.item, quantity=9)
        preview = self.preview()
        self.assertFalse(preview['can_commit'], preview)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 10)

    def test_import_task_ownership_and_revoked_permission(self):
        """A task cannot be committed by another user or after permission loss."""
        preview = self.preview()
        other = User.objects.create_superuser(
            username='other-admin', email='other@example.test', password='test'
        )
        self.client.force_authenticate(other)
        self.post('wms/imports', {'task': preview['task']}, 400)
        self.user.is_superuser = False
        self.user.save()
        self.client.force_authenticate(self.user)
        self.post('wms/imports', {'task': preview['task']}, 403)

    def test_invalid_date_range_and_complete_export(self):
        """Bad dates fail explicitly; pagination never truncates an export."""
        for _ in range(3):
            self.scan()
        self.assertEqual(
            self.client.get(
                '/api/stock/wms/records/', {'start': '2026-09-08', 'end': '2026-09-07'}
            ).status_code,
            400,
        )
        today = timezone.now().date().isoformat()
        response = self.client.get(
            '/api/stock/wms/records/',
            {'start': today, 'end': today, 'limit': 1, 'export': 'csv'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode('utf-8-sig').count('WMS-001'), 3)

    def test_oauth_scopes_apply_to_exports_and_each_import_write(self):
        """Read tokens cannot create previews which authorize writes; scoped writes can."""
        from oauth2_provider.models import AccessToken, Application

        from users.oauth2_scopes import DEFAULT_READ, get_granular_scope

        application = Application.objects.create(
            name='wms-oauth',
            user=self.user,
            client_type='confidential',
            authorization_grant_type='authorization-code',
        )
        token = AccessToken.objects.create(
            user=self.user,
            application=application,
            token='wms-test-token',
            scope=DEFAULT_READ,
            expires=timezone.now() + timedelta(hours=1),
        )
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION='Bearer wms-test-token')
        self.assertEqual(
            self.client.get('/api/stock/wms/inventory/', {'export': 'csv'}).status_code,
            200,
        )
        self.post('wms/imports', {'rows': [self.count_row()]}, 403)
        token.scope = DEFAULT_READ + ' ' + get_granular_scope('change', 'stock')
        token.save()
        preview = self.preview()
        self.assertTrue(preview['can_commit'])
        self.post('wms/imports', {'task': preview['task']})
        token.scope = get_granular_scope('view', 'stock')
        token.save()
        self.assertEqual(self.client.get('/api/stock/wms/inventory/').status_code, 200)


class WmsConcurrencyTest(TransactionTestCase):
    """Use real database connections to race for the last available unit."""

    def test_two_outbound_scans_last_unit(self):
        """PostgreSQL row locks permit at most one competing outbound success."""
        if connection.vendor != 'postgresql':
            self.skipTest('Requires PostgreSQL row locks')
        ContentType.objects.clear_cache()
        user = User.objects.create_superuser(
            username='race-user', email='race@example.test', password='test'
        )
        location = StockLocation.objects.create(name='race-location')
        part = Part.objects.create(name='race-part', IPN='RACE', units='count')
        item = StockItem.objects.create(
            part=part,
            location=location,
            batch='RACE-P',
            quantity=1,
            delete_on_deplete=False,
        )
        set_global_setting('SCAN_APPROVAL_ENABLED', False)
        barrier = threading.Barrier(2)
        results, errors = [], []

        def run():
            try:
                barrier.wait(timeout=10)
                service = ScanStockService(
                    user=user,
                    operation='OUTBOUND',
                    data={
                        'barcode': part.IPN,
                        'quantity': Decimal(1),
                        'batch': 'RACE-P',
                        'purpose': '安装',
                        'source_location': location,
                        'client_event_id': uuid.uuid4(),
                    },
                )
                results.append(service.process().success)
            except Exception as exc:
                errors.append(str(exc))
            finally:
                connections.close_all()

        threads = [threading.Thread(target=run) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=30)
            self.assertFalse(thread.is_alive())
        self.assertEqual(errors, [])
        self.assertEqual(sorted(results), [False, True])
        item.refresh_from_db()
        self.assertEqual(item.quantity, 0)
